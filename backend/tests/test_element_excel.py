"""H 组用例：申报要素 Excel 导入 / 导出。

覆盖：导出内容与筛选、模板下载、xlsx 往返（含其它要素）、导入去重、
无效行报告、表头别名容错、文件类型与角色门槛、未识别表头。
"""
import json
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

HEADERS = ["HS编码", "商品名称", "英文名称", "客户名称", "品牌", "型号", "规格", "材质", "用途", "其它要素"]

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(rows: list[list], headers: list | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers or HEADERS)
    for row in rows:
        ws.append(row)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sheet(data: bytes) -> list[list]:
    """把 xlsx 字节流读成矩阵，便于断言导出内容。"""
    wb = load_workbook(BytesIO(data), data_only=True)
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]


def _upload(client, headers, data: bytes, filename: str = "elements.xlsx"):
    return client.post(
        "/api/elements/import",
        files={"file": (filename, data, XLSX_MIME)},
        headers=headers,
    )


@pytest.fixture()
def seeded(client, users):
    """3 条要素：2 条沃尔玛（8517 系）+ 1 条华为，其中一条带其它要素。"""
    payloads = [
        {"hs_code": "8517.62", "product_name": "蓝牙耳机", "name_en": "BT Earbuds",
         "customer_name": "沃尔玛", "brand": "ONN", "model": "ONN-1001"},
        {"hs_code": "8517.13", "product_name": "手机充电线", "name_en": "Charging Cable",
         "customer_name": "沃尔玛", "brand": "GE"},
        {"hs_code": "8471.30", "product_name": "笔记本电脑", "customer_name": "华为",
         "elements_json": json.dumps([{"name": "成分含量", "value": "ABS 60%"}], ensure_ascii=False)},
    ]
    for p in payloads:
        resp = client.post("/api/elements", json=p, headers=users["customs"]["headers"])
        assert resp.status_code == 200, resp.text
    return payloads


# ---------------- 导出 ----------------

def test_h1_export_requires_login(client):
    assert client.get("/api/elements/export").status_code == 401


def test_h2_export_content_matches_records(client, users, seeded):
    resp = client.get("/api/elements/export", headers=users["customs"]["headers"])
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-type"].startswith(XLSX_MIME)
    # 中文文件名走 RFC 5987 编码
    assert "filename*=UTF-8''" in resp.headers["content-disposition"]

    matrix = _sheet(resp.content)
    assert matrix[0] == HEADERS
    names = {row[1] for row in matrix[1:]}
    assert names == {"蓝牙耳机", "手机充电线", "笔记本电脑"}
    # HS 编码按文本导出，不出现 8517.62 被转成浮点的情况
    assert all(isinstance(row[0], str) for row in matrix[1:])


def test_h3_export_respects_filters(client, users, seeded):
    resp = client.get(
        "/api/elements/export",
        params={"customer_name": "沃尔玛"},
        headers=users["customs"]["headers"],
    )
    matrix = _sheet(resp.content)
    assert len(matrix) == 3  # 表头 + 2 条沃尔玛
    assert {row[3] for row in matrix[1:]} == {"沃尔玛"}


def test_h4_export_serializes_extra_elements(client, users, seeded):
    resp = client.get(
        "/api/elements/export",
        params={"keyword": "笔记本"},
        headers=users["customs"]["headers"],
    )
    matrix = _sheet(resp.content)
    assert matrix[1][9] == "成分含量=ABS 60%"


def test_h5_export_allowed_for_readonly_role(client, users, seeded):
    """导出属只读操作，普通业务员也可用。"""
    assert client.get("/api/elements/export", headers=users["merchandiser"]["headers"]).status_code == 200


# ---------------- 模板 ----------------

def test_h6_template_header_only_and_role_gated(client, users):
    resp = client.get("/api/elements/template", headers=users["customs"]["headers"])
    assert resp.status_code == 200, resp.text
    matrix = _sheet(resp.content)
    assert matrix == [HEADERS]
    assert client.get("/api/elements/template", headers=users["merchandiser"]["headers"]).status_code == 403


# ---------------- 导入 ----------------

def test_h7_import_creates_records(client, users):
    data = _xlsx([
        ["8517.62", "蓝牙耳机", "BT Earbuds", "沃尔玛", "ONN", "ONN-1001", "TWS", "ABS", "听歌", ""],
        ["8471.30", "笔记本电脑", "", "华为", "HUAWEI", "MateBook D14", "14英寸", "", "", ""],
    ])
    body = _upload(client, users["customs"]["headers"], data).json()
    assert body == {"total": 2, "created": 2, "skipped": 0, "invalid": 0, "errors": []}

    listed = client.get("/api/elements", headers=users["customs"]["headers"]).json()
    assert listed["total"] == 2
    assert {r["product_name"] for r in listed["items"]} == {"蓝牙耳机", "笔记本电脑"}


def test_h8_import_skips_duplicates(client, users, seeded):
    """重复导入同一文件：第二次全部按重复跳过，不新增也不覆盖。"""
    resp = client.get("/api/elements/export", headers=users["customs"]["headers"])
    body = _upload(client, users["customs"]["headers"], resp.content).json()
    assert body["created"] == 0
    assert body["skipped"] == 3
    assert client.get("/api/elements", headers=users["customs"]["headers"]).json()["total"] == 3


def test_h9_import_reports_invalid_rows_with_line_number(client, users):
    data = _xlsx([
        ["8517.62", "蓝牙耳机", "", "沃尔玛", "", "", "", "", "", ""],
        ["", "缺编码的商品", "", "沃尔玛", "", "", "", "", "", ""],
        ["8517.13", "手机充电线", "", "", "", "", "", "", "", ""],
    ])
    body = _upload(client, users["customs"]["headers"], data).json()
    assert body["total"] == 3
    assert body["created"] == 1
    assert body["invalid"] == 2
    # 行号对应用户在 Excel 中看到的位置（表头为第 1 行）
    assert body["errors"][0] == "第 3 行：缺少HS编码"
    assert body["errors"][1] == "第 4 行：缺少客户名称"


def test_h10_import_accepts_header_aliases(client, users):
    data = _xlsx(
        [["8517.62", "蓝牙耳机", "沃尔玛"]],
        headers=["HS编码", "品名", "客户"],
    )
    body = _upload(client, users["customs"]["headers"], data).json()
    assert body["created"] == 1
    item = client.get("/api/elements", headers=users["customs"]["headers"]).json()["items"][0]
    assert item["product_name"] == "蓝牙耳机"
    assert item["customer_name"] == "沃尔玛"


def test_h11_import_round_trips_extra_elements(client, users, seeded):
    resp = client.get("/api/elements/export", headers=users["customs"]["headers"])
    matrix = _sheet(resp.content)
    # 换一个客户名，避免命中去重
    for row in matrix[1:]:
        row[3] = f"{row[3]}-2"
    body = _upload(client, users["customs"]["headers"], _xlsx(matrix[1:])).json()
    assert body["created"] == 3

    listed = client.get(
        "/api/elements", params={"customer_name": "华为-2"}, headers=users["customs"]["headers"]
    ).json()
    extra = json.loads(listed["items"][0]["elements_json"])
    assert extra == [{"name": "成分含量", "value": "ABS 60%"}]


def test_h12_import_rejects_non_xlsx(client, users):
    resp = _upload(client, users["customs"]["headers"], b"a,b\n1,2\n", filename="elements.csv")
    assert resp.status_code == 400
    assert "xlsx" in resp.json()["detail"]
    resp = _upload(client, users["customs"]["headers"], b"", filename="elements.xlsx")
    assert resp.status_code == 400


def test_h13_import_requires_customs_role(client, users):
    data = _xlsx([["8517.62", "蓝牙耳机", "", "沃尔玛", "", "", "", "", "", ""]])
    assert _upload(client, users["merchandiser"]["headers"], data).status_code == 403


def test_h14_import_unrecognized_header(client, users):
    data = _xlsx([["a", "b"]], headers=["列一", "列二"])
    body = _upload(client, users["customs"]["headers"], data).json()
    assert body["created"] == 0
    assert body["invalid"] == 1
    assert "未识别到表头" in body["errors"][0]
