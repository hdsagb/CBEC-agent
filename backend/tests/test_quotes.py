"""E 组用例：智能比价 F1 —— 上传 → 解析 → 确认 → 加权比价。

覆盖：Excel 表头规则映射、Word/PDF 无 Key 降级、重复解析幂等、确认流、
权重对最优供应商的影响、未确认明细不参与比价、权限与文件类型校验。
"""
import io

import pytest
from docx import Document
from openpyxl import Workbook

from app.parsers import router as parsers

SUPPLIER_A_ROWS = [
    ["产品名称", "规格", "单价", "币种", "最小起订量", "交货期", "有效期", "备注"],
    ["蓝牙耳机", "TWS", "3.5", "USD", "1000", "20", "2026-12-31", "A厂现货"],
    ["充电线", "USB-C 1m", "1.2", "USD", "2000", "25", "2026-12-31", ""],
]

SUPPLIER_B_ROWS = [
    ["产品名称", "规格", "单价", "最小起订量", "交货期"],
    ["蓝牙耳机", "TWS", "3.2", "500", "35"],
    ["充电线", "USB-C 1m", "1.5", "1000", "20"],
]


def _xlsx_bytes(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(rows: list[list]) -> bytes:
    doc = Document()
    table = doc.add_table(rows=0, cols=len(rows[0]))
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture()
def batch_id(client, users) -> int:
    resp = client.post(
        "/api/quotes/batches", params={"name": "Q3 采购比价"}, headers=users["merchandiser"]["headers"]
    )
    assert resp.status_code == 200, resp.text
    return resp.json()["id"]


def _upload(client, headers, batch_id: int, filename: str, content: bytes, supplier: str = ""):
    resp = client.post(
        f"/api/quotes/batches/{batch_id}/sources",
        params={"supplier_name": supplier},
        files={"file": (filename, content, "application/octet-stream")},
        headers=headers,
    )
    assert resp.status_code == 200, resp.text
    return resp.json()


def _parse_and_confirm(client, headers, batch_id: int) -> list[dict]:
    parsed = client.post(f"/api/quotes/batches/{batch_id}/parse", headers=headers)
    assert parsed.status_code == 200, parsed.text
    items = client.get(f"/api/quotes/batches/{batch_id}/items", headers=headers).json()
    for item in items:
        resp = client.put(
            f"/api/quotes/items/{item['id']}/confirm",
            json={
                "product_name": item["product_name"],
                "spec": item["spec"],
                "unit_price": item["unit_price"],
                "currency": item["currency"],
                "moq": item["moq"],
                "lead_time_days": item["lead_time_days"],
                "valid_until": item["valid_until"],
                "remarks": item["remarks"],
            },
            headers=headers,
        )
        assert resp.status_code == 200, resp.text
    return items


def test_e1_create_and_upload_xlsx(client, users, batch_id):
    headers = users["merchandiser"]["headers"]
    body = _upload(client, headers, batch_id, "supplier_a.xlsx", _xlsx_bytes(SUPPLIER_A_ROWS), "供应商A")
    assert len(body["sources"]) == 1
    source = body["sources"][0]
    assert source["file_type"] == "xlsx"
    assert source["parse_status"] == "pending"
    assert source["supplier_name"] == "供应商A"
    assert body["status"] == "parsing"


def test_e2_excel_header_mapping(client, users, batch_id):
    headers = users["merchandiser"]["headers"]
    _upload(client, headers, batch_id, "supplier_a.xlsx", _xlsx_bytes(SUPPLIER_A_ROWS), "供应商A")
    resp = client.post(f"/api/quotes/batches/{batch_id}/parse", headers=headers)
    assert resp.status_code == 200, resp.text
    items = resp.json()
    assert len(items) == 2

    earbuds = next(i for i in items if i["product_name"] == "蓝牙耳机")
    assert earbuds["spec"] == "TWS"
    assert earbuds["unit_price"] == 3.5
    assert earbuds["currency"] == "USD"
    assert earbuds["moq"] == 1000
    assert earbuds["lead_time_days"] == 20
    assert earbuds["valid_until"] == "2026-12-31"
    assert earbuds["remarks"] == "A厂现货"
    assert earbuds["confidence"] > 0.6
    assert earbuds["confirmed"] is False


def test_e3_reparse_is_idempotent(client, users, batch_id):
    headers = users["merchandiser"]["headers"]
    _upload(client, headers, batch_id, "supplier_a.xlsx", _xlsx_bytes(SUPPLIER_A_ROWS), "供应商A")
    client.post(f"/api/quotes/batches/{batch_id}/parse", headers=headers)
    client.post(f"/api/quotes/batches/{batch_id}/parse", headers=headers)
    items = client.get(f"/api/quotes/batches/{batch_id}/items", headers=headers).json()
    assert len(items) == 2  # 重复解析不叠加


def test_e4_word_fallback_without_api_key(client, users, batch_id, monkeypatch):
    """无 DEEPSEEK_API_KEY 时 Word 走规则降级，仍能解析出明细。"""
    monkeypatch.setattr(parsers, "llm_extract", lambda _text: None)
    headers = users["merchandiser"]["headers"]
    _upload(client, headers, batch_id, "supplier_b.docx", _docx_bytes(SUPPLIER_B_ROWS), "供应商B")
    resp = client.post(f"/api/quotes/batches/{batch_id}/parse", headers=headers)
    assert resp.status_code == 200, resp.text
    items = resp.json()
    assert len(items) == 2
    earbuds = next(i for i in items if i["product_name"] == "蓝牙耳机")
    assert earbuds["unit_price"] == 3.2
    assert earbuds["moq"] == 500
    assert earbuds["lead_time_days"] == 35


def test_e5_pdf_fallback_without_api_key(client, users, batch_id, monkeypatch):
    """PDF 无 Key 时按行分列 + 表头映射降级解析。"""
    monkeypatch.setattr(parsers, "llm_extract", lambda _text: None)
    monkeypatch.setattr(
        parsers,
        "pdf_text",
        lambda _path: "产品名称    规格    单价    最小起订量    交货期\n蓝牙耳机    TWS    3.9    800    28\n",
    )
    headers = users["merchandiser"]["headers"]
    _upload(client, headers, batch_id, "supplier_c.pdf", b"%PDF-1.4 dummy", "供应商C")
    resp = client.post(f"/api/quotes/batches/{batch_id}/parse", headers=headers)
    assert resp.status_code == 200, resp.text
    items = resp.json()
    assert len(items) == 1
    assert items[0]["product_name"] == "蓝牙耳机"
    assert items[0]["unit_price"] == 3.9
    assert items[0]["lead_time_days"] == 28


def test_e6_confirm_advances_batch_status(client, users, batch_id):
    headers = users["merchandiser"]["headers"]
    _upload(client, headers, batch_id, "supplier_a.xlsx", _xlsx_bytes(SUPPLIER_A_ROWS), "供应商A")
    _parse_and_confirm(client, headers, batch_id)
    batch = client.get(f"/api/quotes/batches/{batch_id}", headers=headers).json()
    assert batch["status"] == "confirmed"


def test_e7_compare_weight_changes_best_supplier(client, users, batch_id):
    headers = users["merchandiser"]["headers"]
    _upload(client, headers, batch_id, "supplier_a.xlsx", _xlsx_bytes(SUPPLIER_A_ROWS), "供应商A")
    _upload(client, headers, batch_id, "supplier_b.xlsx", _xlsx_bytes(SUPPLIER_B_ROWS), "供应商B")
    _parse_and_confirm(client, headers, batch_id)

    # 价格优先：两家都应由更便宜的供应商胜出
    by_price = client.post(
        f"/api/quotes/batches/{batch_id}/compare", json={"weight_config": {"price": 1}}, headers=headers
    ).json()
    assert by_price["weights"]["price"] == 1.0
    best_by_price = {g["product_name"]: g["best_supplier"] for g in by_price["result"]}
    assert best_by_price["蓝牙耳机"] == "供应商B"  # 3.2 < 3.5
    assert best_by_price["充电线"] == "供应商A"  # 1.2 < 1.5

    # 货期优先：交期更短的供应商胜出
    by_lead = client.post(
        f"/api/quotes/batches/{batch_id}/compare", json={"weight_config": {"lead_time": 1}}, headers=headers
    ).json()
    best_by_lead = {g["product_name"]: g["best_supplier"] for g in by_lead["result"]}
    assert best_by_lead["蓝牙耳机"] == "供应商A"  # 20 < 35
    assert best_by_lead["充电线"] == "供应商B"  # 20 < 25

    # 批次状态推进到 compared，且候选带四维得分
    batch = client.get(f"/api/quotes/batches/{batch_id}", headers=headers).json()
    assert batch["status"] == "compared"
    candidate = by_price["result"][0]["candidates"][0]
    assert set(candidate["scores"]) == {"price", "moq", "lead_time", "confidence"}
    assert candidate["is_best"] is True


def test_e8_unconfirmed_items_excluded(client, users, batch_id):
    headers = users["merchandiser"]["headers"]
    _upload(client, headers, batch_id, "supplier_a.xlsx", _xlsx_bytes(SUPPLIER_A_ROWS), "供应商A")
    client.post(f"/api/quotes/batches/{batch_id}/parse", headers=headers)
    result = client.post(f"/api/quotes/batches/{batch_id}/compare", json={}, headers=headers).json()
    assert result["result"] == []
    assert result["recommended_source_id"] is None


def test_e9_unsupported_file_type_rejected(client, users, batch_id):
    resp = client.post(
        f"/api/quotes/batches/{batch_id}/sources",
        files={"file": ("evil.exe", b"x", "application/octet-stream")},
        headers=users["merchandiser"]["headers"],
    )
    assert resp.status_code == 400


def test_e10_requires_auth(client, batch_id):
    assert client.post("/api/quotes/batches", params={"name": "x"}).status_code == 401
    assert client.get(f"/api/quotes/batches/{batch_id}").status_code == 401


def test_e11_batch_list_pagination(client, users):
    """批次列表按创建时间倒序分页，翻页结果不重叠且 total 为全量条数。"""
    headers = users["merchandiser"]["headers"]
    created = []
    for i in range(5):
        resp = client.post(
            "/api/quotes/batches", params={"name": f"分页批次{i}"}, headers=headers
        )
        assert resp.status_code == 200, resp.text
        created.append(resp.json()["id"])

    first = client.get(
        "/api/quotes/batches", params={"page": 1, "page_size": 2}, headers=headers
    ).json()
    assert first["total"] == 5
    assert first["page"] == 1 and first["page_size"] == 2
    # 最新创建的批次排在第一页最前
    assert [b["id"] for b in first["items"]] == [created[-1], created[-2]]

    second = client.get(
        "/api/quotes/batches", params={"page": 2, "page_size": 2}, headers=headers
    ).json()
    assert [b["id"] for b in second["items"]] == [created[-3], created[-4]]

    third = client.get(
        "/api/quotes/batches", params={"page": 3, "page_size": 2}, headers=headers
    ).json()
    assert [b["id"] for b in third["items"]] == [created[0]]

    assert client.get(
        "/api/quotes/batches", params={"page": 0}, headers=headers
    ).status_code == 422
    assert client.get(
        "/api/quotes/batches", params={"page_size": 101}, headers=headers
    ).status_code == 422
