"""申报要素库 Excel 导入 / 导出。

设计要点：
- 导出：按固定列序写入 xlsx；「其它要素」序列化为 `名称=值; 名称=值` 单列，
  导入时按同样格式还原，保证导出→再导入可无损往返。
- 导入：读取首个工作表，前几行内定位表头（归一化后先精确匹配再子串匹配别名），
  逐行校验必填项，返回标准化记录与逐行错误说明，不在此处访问数据库。
- HS 编码等字段一律按文本写入，避免被 Excel 识别为数值造成精度丢失。
"""

import json
import re
from io import BytesIO

# 导出列序：(字段名, 中文表头)
EXPORT_COLUMNS: list[tuple[str, str]] = [
    ("hs_code", "HS编码"),
    ("product_name", "商品名称"),
    ("name_en", "英文名称"),
    ("customer_name", "客户名称"),
    ("brand", "品牌"),
    ("model", "型号"),
    ("spec", "规格"),
    ("material", "材质"),
    ("usage", "用途"),
    ("elements_json", "其它要素"),
]

# 必填字段与中文名（缺一即视为无效行）
REQUIRED_FIELDS: dict[str, str] = {
    "hs_code": "HS编码",
    "product_name": "商品名称",
    "customer_name": "客户名称",
}

# 表头别名（均为归一化后的形式）
FIELD_ALIASES: dict[str, list[str]] = {
    "hs_code": ["hs编码", "hs", "编码", "商品编码", "税则号列", "税则号"],
    "product_name": ["商品名称", "产品名称", "品名", "中文品名", "货物名称"],
    "name_en": ["英文名称", "英文品名", "英文名"],
    "customer_name": ["客户名称", "客户", "国外客户", "收货人"],
    "brand": ["品牌", "牌子"],
    "model": ["型号", "规格型号"],
    "spec": ["规格", "规格描述"],
    "material": ["材质", "成分", "材料"],
    "usage": ["用途", "功能"],
    "elements_json": ["其它要素", "其他要素", "其它申报要素", "附加要素"],
}

_EXTRA_PAIR_SEP = ";"
_EXTRA_KV_SEP = "="
_HEADER_SCAN_ROWS = 5
_MIN_HEADER_FIELDS = 2


def _norm(text: object) -> str:
    """归一化表头：小写、去除空白与常见标点，便于别名匹配。"""
    s = str(text or "").strip().lower()
    return re.sub(r"[\s:：\-_/\\()（）\[\]【】.。、,，*]+", "", s)


_ALIAS_INDEX: list[tuple[str, str]] = sorted(
    ((_norm(a), f) for f, aliases in FIELD_ALIASES.items() for a in aliases),
    key=lambda x: len(x[0]),
    reverse=True,
)


def _match_field(header: object) -> str | None:
    h = _norm(header)
    if not h:
        return None
    for alias, field in _ALIAS_INDEX:
        if h == alias:
            return field
    for alias, field in _ALIAS_INDEX:
        if len(alias) >= 2 and alias in h:
            return field
    return None


def _str(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # Excel 会把 8517 读成 8517.0，去掉多余小数位
        return str(int(value))
    return str(value).strip()


def _parse_extra(text: object) -> str:
    """`名称=值; 名称=值` → elements_json 字符串。"""
    raw = _str(text)
    if not raw:
        return "[]"
    items: list[dict] = []
    for chunk in raw.split(_EXTRA_PAIR_SEP):
        chunk = chunk.strip()
        if not chunk:
            continue
        name, _, value = chunk.partition(_EXTRA_KV_SEP)
        name, value = name.strip(), value.strip()
        if name or value:
            items.append({"name": name, "value": value})
    return json.dumps(items, ensure_ascii=False)


def _dump_extra(elements_json: str) -> str:
    """elements_json 字符串 → `名称=值; 名称=值`。"""
    try:
        data = json.loads(elements_json or "[]")
    except (TypeError, ValueError):
        return ""
    if not isinstance(data, list):
        return ""
    pairs: list[str] = []
    for item in data:
        if not isinstance(item, dict):
            continue
        name = _str(item.get("name"))
        value = _str(item.get("value"))
        if name or value:
            pairs.append(f"{name}{_EXTRA_KV_SEP}{value}")
    return f"{_EXTRA_PAIR_SEP} ".join(pairs)


def _cell(element: object, field: str) -> str:
    if field == "elements_json":
        return _dump_extra(getattr(element, "elements_json", "") or "")
    return _str(getattr(element, field, ""))


def build_workbook(elements: list | None = None) -> bytes:
    """把要素记录（可为空，用于生成导入模板）写成 xlsx 字节流。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "申报要素"

    ws.append([label for _, label in EXPORT_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for element in elements or []:
        ws.append([_cell(element, field) for field, _ in EXPORT_COLUMNS])

    widths = [14, 26, 20, 18, 14, 16, 18, 14, 20, 30]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_workbook(data: bytes) -> tuple[list[dict], list[str]]:
    """解析 xlsx 字节流。返回 (标准化记录, 逐行错误说明)。"""
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    try:
        matrix = [list(row) for row in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()

    header_idx, col_map = _locate_header(matrix)
    if header_idx < 0:
        return [], ["未识别到表头，请确认包含「HS编码 / 商品名称 / 客户名称」等列"]

    records: list[dict] = []
    errors: list[str] = []
    for offset, row in enumerate(matrix[header_idx + 1:], start=header_idx + 2):
        if not row or all(_str(c) == "" for c in row):
            continue
        raw = {field: "" for field, _ in EXPORT_COLUMNS}
        for col, field in col_map.items():
            if col < len(row):
                raw[field] = row[col]
        rec = {
            "hs_code": _str(raw["hs_code"]),
            "product_name": _str(raw["product_name"]),
            "name_en": _str(raw["name_en"]),
            "customer_name": _str(raw["customer_name"]),
            "brand": _str(raw["brand"]),
            "model": _str(raw["model"]),
            "spec": _str(raw["spec"]),
            "material": _str(raw["material"]),
            "usage": _str(raw["usage"]),
            "elements_json": _parse_extra(raw["elements_json"]),
        }
        missing = [cn for f, cn in REQUIRED_FIELDS.items() if not rec[f]]
        if missing:
            errors.append(f"第 {offset} 行：缺少{'、'.join(missing)}")
            continue
        records.append(rec)
    return records, errors


def _locate_header(matrix: list[list]) -> tuple[int, dict[int, str]]:
    """在前几行内定位表头行，返回 (行下标, {列下标: 字段名})；未识别返回 (-1, {})。"""
    best_idx = -1
    best_map: dict[int, str] = {}
    for i, row in enumerate(matrix[:_HEADER_SCAN_ROWS]):
        mapping: dict[int, str] = {}
        for j, cell in enumerate(row):
            field = _match_field(cell)
            if field and field not in mapping.values():
                mapping[j] = field
        if len(mapping) > len(best_map):
            best_map, best_idx = mapping, i
    if len(best_map) < _MIN_HEADER_FIELDS:
        return -1, {}
    return best_idx, best_map
