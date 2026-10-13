"""报价单解析管道：格式路由 → 规则抽取 →（可选）LLM 结构化增强。

设计要点：
- Excel：按表头别名规则映射为标准化行（不依赖 LLM）。
- Word / PDF：优先调用 DeepSeek JSON Mode 抽取；未配置 API Key 或调用失败时
  自动降级为规则抽取（Word 表格按表头映射、PDF 按行分列后按表头映射），
  保证端到端可跑通、可测试。
- 返回标准化行：{product_name, spec, unit_price, currency, moq, lead_time_days, valid_until, remarks}
"""

import json
import re

CANONICAL = [
    "product_name",
    "spec",
    "unit_price",
    "currency",
    "moq",
    "lead_time_days",
    "valid_until",
    "remarks",
]

# 表头别名表：主字段 → 常见中英文表头（会先做归一化再去匹配）
FIELD_ALIASES: dict[str, list[str]] = {
    "product_name": [
        "产品名称", "商品名称", "货物名称", "品名", "产品", "商品",
        "productname", "product", "description", "itemname",
    ],
    "spec": [
        "规格型号", "规格", "型号", "spec", "specification", "model", "type",
    ],
    "unit_price": [
        "含税单价", "不含税单价", "单价", "价格", "报价",
        "unitprice", "price", "fob",
    ],
    "currency": ["币种", "货币", "结算币种", "currency", "curr"],
    "moq": [
        "最小起订量", "起订量", "最小订量", "moq",
        "minorder", "minimumorder", "minorderqty",
    ],
    "lead_time_days": [
        "交货期", "生产周期", "货期", "交期",
        "leadtimedays", "leadtimeday", "leadtime", "deliverytime", "delivery",
    ],
    "valid_until": ["报价有效期", "有效期", "validuntil", "validity", "validto"],
    "remarks": ["备注", "说明", "remark", "remarks", "note", "comment", "comments"],
}

EXTRACTION_PROMPT = """你是报价单结构化助手。从以下表格文本中抽取报价明细，
以 JSON 输出 {"items": [...]}，每项包含：
product_name(产品名), spec(规格), unit_price(单价数字), currency(币种),
moq(最小起订量), lead_time_days(货期天数), valid_until(有效期), remarks(备注)。
对不确定的字段填 null。"""


def _norm(text: object) -> str:
    """归一化表头：小写、去除空白与常见标点，便于别名匹配。"""
    s = str(text or "").strip().lower()
    return re.sub(r"[\s:：\-_/\\()（）\[\]【】.。、,，*]+", "", s)


_ALIAS_INDEX: list[tuple[str, str]] = sorted(
    ((_norm(a), f) for f, aliases in FIELD_ALIASES.items() for a in aliases),
    key=lambda x: len(x[0]),
    reverse=True,
)


def _match_field(header: object) -> str | None:
    h = _norm(header)
    if not h:
        return None
    for alias, field in _ALIAS_INDEX:
        if h == alias:
            return field
    for alias, field in _ALIAS_INDEX:
        if len(alias) >= 2 and alias in h:
            return field
    return None


def _str(value: object) -> str:
    return "" if value is None else str(value).strip()


def _to_float(value: object) -> float | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    m = re.search(r"-?\d+(?:\.\d+)?", str(value).replace(",", ""))
    return float(m.group()) if m else None


def _to_int(value: object) -> int | None:
    f = _to_float(value)
    return int(f) if f is not None else None


_CURRENCY_HINTS: list[tuple[str, str]] = [
    ("usd", "USD"), ("$", "USD"),
    ("eur", "EUR"), ("€", "EUR"),
    ("gbp", "GBP"), ("£", "GBP"),
    ("rmb", "CNY"), ("cny", "CNY"), ("￥", "CNY"), ("¥", "CNY"), ("元", "CNY"),
]


def _detect_currency(*texts: object, default: str = "CNY") -> str:
    joined = " ".join(str(t) for t in texts if t).lower()
    for hint, code in _CURRENCY_HINTS:
        if hint in joined:
            return code
    return default


def _coerce(raw: dict) -> dict:
    """把任意来源的原始行规整为标准化行（兼容字符串/数字/None）。"""
    price_text = raw.get("unit_price")
    return {
        "product_name": _str(raw.get("product_name")),
        "spec": _str(raw.get("spec")),
        "unit_price": _to_float(price_text),
        "currency": _detect_currency(raw.get("currency"), price_text),
        "moq": _to_int(raw.get("moq")),
        "lead_time_days": _to_int(raw.get("lead_time_days")),
        "valid_until": _str(raw.get("valid_until")),
        "remarks": _str(raw.get("remarks")),
    }


def _split_line_columns(line: str) -> list[str]:
    """PDF 文本按 2+ 空白或制表符切列。"""
    return [c.strip() for c in re.split(r"\s{2,}|\t+", line.strip()) if c.strip()]


def rows_from_matrix(matrix: list[list]) -> list[dict]:
    """在表格矩阵中定位表头行并按别名映射出标准化行。"""
    if not matrix:
        return []

    best_idx = -1
    best_map: dict[int, str] = {}
    for i, row in enumerate(matrix[:10]):
        mapping: dict[int, str] = {}
        for j, cell in enumerate(row):
            field = _match_field(cell)
            if field and field not in mapping.values():
                mapping[j] = field
        if len(mapping) > len(best_map):
            best_map, best_idx = mapping, i

    # 至少识别出 2 个字段才认为存在有效表头
    if len(best_map) < 2:
        return []

    rows: list[dict] = []
    for row in matrix[best_idx + 1:]:
        if not row or all(_str(c) == "" for c in row):
            continue
        raw = {f: "" for f in CANONICAL}
        for j, field in best_map.items():
            if j < len(row):
                raw[field] = row[j]
        rec = _coerce(raw)
        if not rec["product_name"] and rec["unit_price"] is None:
            continue
        rows.append(rec)
    return rows


def extract_excel(file_path: str) -> tuple[list[dict], float]:
    """Excel：逐表按表头规则映射。返回 (标准化行, 基础置信度)。"""
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    rows: list[dict] = []
    for sheet in wb.worksheets:
        matrix = [list(r) for r in sheet.iter_rows(values_only=True)]
        rows.extend(rows_from_matrix(matrix))
    return rows, 0.9


def word_text(file_path: str) -> str:
    from docx import Document

    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def extract_word(file_path: str) -> tuple[list[dict], float]:
    """Word 规则降级：读取表格并按表头映射。"""
    from docx import Document

    doc = Document(file_path)
    rows: list[dict] = []
    for table in doc.tables:
        matrix = [[cell.text for cell in row.cells] for row in table.rows]
        rows.extend(rows_from_matrix(matrix))
    return rows, 0.85


def pdf_text(file_path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def extract_pdf(file_path: str) -> tuple[list[dict], float]:
    """PDF 规则降级：抽取文本 → 按行分列 → 复用表头映射。"""
    text = pdf_text(file_path)
    matrix = [_split_line_columns(line) for line in text.splitlines()]
    return rows_from_matrix(matrix), 0.6


def read_text(file_path: str, file_type: str) -> str:
    """读取供 LLM 抽取的原始文本。"""
    try:
        if file_type == "docx":
            return word_text(file_path)
        if file_type == "pdf":
            return pdf_text(file_path)
        if file_type in ("xlsx", "xls"):
            from openpyxl import load_workbook

            wb = load_workbook(file_path, read_only=True, data_only=True)
            lines: list[str] = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    if row and any(c is not None for c in row):
                        lines.append(" | ".join(_str(c) for c in row))
            return "\n".join(lines)
        if file_type in IMAGE_TYPES:
            return ocr_text(file_path)
    except Exception:
        return ""
    return ""


def _items_from_json(data: object) -> list[dict] | None:
    """把 LLM 返回的 {"items": [...]} 规整为标准化行；结构不合法返回 None。"""
    if not isinstance(data, dict):
        return None
    items = data.get("items")
    if not isinstance(items, list):
        return None
    out: list[dict] = []
    for item in items:
        if isinstance(item, dict):
            out.append(_coerce(item))
    return out or None


def llm_extract(raw_text: str) -> list[dict] | None:
    """调用 LLM JSON Mode 抽取；不可用（无 Key / 网络失败 / 解析失败）返回 None。"""
    if not raw_text.strip():
        return None
    try:
        from app.services.llm.base import get_llm_client

        client = get_llm_client()
        resp = client.chat(
            [
                {"role": "system", "content": EXTRACTION_PROMPT},
                {"role": "user", "content": raw_text[:8000]},
            ],
            json_mode=True,
        )
        data = json.loads(resp.get("content") or "{}")
    except Exception:
        return None
    return _items_from_json(data)


# ---------------- 图片报价单：多模态 LLM 优先，本地 OCR 兜底 ----------------

IMAGE_TYPES = ("png", "jpg", "jpeg")

_IMAGE_MIME = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg"}

IMAGE_PROMPT = "请识别这张报价单图片，抽取其中的报价明细并结构化输出。"


def image_data_url(file_path: str, file_type: str) -> str:
    """把图片读成 base64 data URL，供多模态模型使用。"""
    import base64

    mime = _IMAGE_MIME.get(file_type.lower().lstrip("."), "image/png")
    with open(file_path, "rb") as fh:
        payload = base64.b64encode(fh.read()).decode("ascii")
    return f"data:{mime};base64,{payload}"


def llm_extract_image(file_path: str, file_type: str) -> list[dict] | None:
    """多模态 LLM 抽取图片报价单；未配置视觉模型或调用失败时返回 None。"""
    try:
        from app.services.llm.base import get_vision_client

        client = get_vision_client()
        if client is None:
            return None
        resp = client.chat(
            [
                {"role": "system", "content": EXTRACTION_PROMPT},
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": IMAGE_PROMPT},
                        {
                            "type": "image_url",
                            "image_url": {"url": image_data_url(file_path, file_type)},
                        },
                    ],
                },
            ],
            json_mode=True,
        )
        data = json.loads(resp.get("content") or "{}")
    except Exception:
        return None
    return _items_from_json(data)


_OCR_ENGINE = None
_OCR_UNAVAILABLE = False


def _get_ocr_engine():
    """懒加载 RapidOCR 引擎；OCR 关闭、依赖缺失或初始化失败时返回 None。"""
    global _OCR_ENGINE, _OCR_UNAVAILABLE
    if _OCR_ENGINE is not None or _OCR_UNAVAILABLE:
        return _OCR_ENGINE
    from app.core.config import settings

    if not settings.ocr_enabled:
        _OCR_UNAVAILABLE = True
        return None
    try:
        from rapidocr import RapidOCR

        _OCR_ENGINE = RapidOCR()
    except Exception:
        _OCR_UNAVAILABLE = True
    return _OCR_ENGINE


def _box_to_line(box: object, text: object) -> tuple[float, float, str]:
    """把 OCR 文字框转成 (纵向中心 y, 最左 x, 文本)。"""
    try:
        ys = [float(p[1]) for p in box]  # type: ignore[index]
        xs = [float(p[0]) for p in box]  # type: ignore[index]
        return (sum(ys) / len(ys), min(xs), str(text))
    except Exception:
        return (0.0, 0.0, str(text))


def ocr_lines(file_path: str) -> list[tuple[float, float, str]]:
    """本地 OCR：返回 (y_center, x_left, text)；引擎不可用时返回空列表。"""
    engine = _get_ocr_engine()
    if engine is None:
        return []
    try:
        raw = engine(file_path)
    except Exception:
        return []

    lines: list[tuple[float, float, str]] = []
    # rapidocr 2.x：返回带 boxes / txts 属性的结果对象
    boxes = getattr(raw, "boxes", None)
    txts = getattr(raw, "txts", None)
    if boxes is not None and txts is not None:
        for box, text in zip(boxes, txts):
            lines.append(_box_to_line(box, text))
        return lines

    # rapidocr_onnxruntime 1.x：返回 ([[box, text, score], ...], elapse)
    data = raw[0] if isinstance(raw, tuple) and raw else raw
    if isinstance(data, (list, tuple)):
        for entry in data:
            if isinstance(entry, (list, tuple)) and len(entry) >= 2:
                lines.append(_box_to_line(entry[0], entry[1]))
    return lines


def ocr_text(file_path: str) -> str:
    """本地 OCR 文本（按纵向、横向排序），供 read_text 复用。"""
    lines = ocr_lines(file_path)
    return "\n".join(text for _, _, text in sorted(lines, key=lambda t: (t[0], t[1])))


def _matrix_from_ocr(lines: list[tuple[float, float, str]]) -> list[list[str]]:
    """把 OCR 文字框按 y 聚类成行、行内按 x 排序，重建表格矩阵。"""
    if not lines:
        return []
    ordered = sorted(lines, key=lambda t: (t[0], t[1]))
    gaps = sorted(
        ordered[i + 1][0] - ordered[i][0]
        for i in range(len(ordered) - 1)
        if ordered[i + 1][0] - ordered[i][0] > 0.5
    )
    tol = (gaps[len(gaps) // 2] if gaps else 12.0) * 0.7

    matrix: list[list[str]] = []
    current: list[tuple[float, str]] = []
    base_y = ordered[0][0]
    for y, x, text in ordered:
        if current and abs(y - base_y) > tol:
            matrix.append([t for _, t in sorted(current)])
            current = []
        if not current:
            base_y = y
        current.append((x, text))
    if current:
        matrix.append([t for _, t in sorted(current)])
    return matrix


def extract_image(file_path: str, file_type: str) -> tuple[list[dict], float]:
    """图片报价单识别：多模态 LLM → 本地 OCR 文本 → 表头映射。"""
    rows = llm_extract_image(file_path, file_type)
    if rows:
        return rows, 0.8
    rows = rows_from_matrix(_matrix_from_ocr(ocr_lines(file_path)))
    return rows, 0.6


def row_confidence(base: float, rec: dict) -> float:
    """按关键字段完整度微调单行置信度，供前端徽章展示。"""
    score = base
    if not rec.get("product_name"):
        score *= 0.7
    if rec.get("unit_price") is None:
        score *= 0.7
    return round(min(max(score, 0.1), 1.0), 2)